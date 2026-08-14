from __future__ import annotations

"""Dedicated parameter-memory operations for isolated V13 calibration."""

from datetime import datetime
from pathlib import Path
from typing import Any
import shutil

import pandas as pd
from openpyxl import load_workbook


BEST_FILE_NAME = "best_parameters_V13.xlsx"


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def best_file(results_dir: str | Path) -> Path:
    return Path(results_dir) / BEST_FILE_NAME


def has_best(results_dir: str | Path) -> bool:
    return best_file(results_dir).exists()


def save_best(
    config_file: str | Path,
    results_dir: str | Path,
    *,
    total_score: float,
    run_folder: str,
    objective_reference_file: str = "",
    objective_mode: str = "multi_species_frozen_reference_weighted_rmse",
    optimizer_version: str = "V13.2",
) -> dict[str, Any]:
    """Copy active configuration to V13 best memory and append metadata."""
    config_file = Path(config_file)
    output = best_file(results_dir)
    output.parent.mkdir(parents=True, exist_ok=True)
    if not config_file.exists():
        raise FileNotFoundError(f"V13 config file does not exist: {config_file}")

    if output.exists():
        backup = output.with_name(f"{output.stem}_previous_{datetime.now():%Y%m%d_%H%M%S}{output.suffix}")
        shutil.copy2(output, backup)
    else:
        backup = None

    shutil.copy2(config_file, output)

    metadata = {
        "created_at": _now(),
        "source_run_folder": str(run_folder),
        "TOTAL_SCORE": float(total_score),
        "objective_reference_file": str(objective_reference_file),
        "objective_mode": str(objective_mode),
        "optimizer_version": str(optimizer_version),
    }

    workbook = load_workbook(output)
    if "v13_metadata" in workbook.sheetnames:
        del workbook["v13_metadata"]
    ws = workbook.create_sheet("v13_metadata")
    ws.append(list(metadata.keys()))
    ws.append(list(metadata.values()))
    workbook.save(output)

    return {
        "action": "saved_v13_best",
        "best_file": str(output),
        "backup_file": str(backup) if backup else "",
        **metadata,
    }


def restore_best(config_file: str | Path, results_dir: str | Path) -> dict[str, Any]:
    """Restore active agent_config from the dedicated V13 best file."""
    config_file = Path(config_file)
    source = best_file(results_dir)
    if not source.exists():
        raise FileNotFoundError(f"V13 best parameter file does not exist: {source}")

    backup_dir = Path(results_dir) / "v13_parameter_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"{config_file.stem}_before_v13_restore_{datetime.now():%Y%m%d_%H%M%S}{config_file.suffix}"
    if config_file.exists():
        shutil.copy2(config_file, backup)

    shutil.copy2(source, config_file)
    return {
        "action": "restored_v13_best",
        "best_file": str(source),
        "config_backup_file": str(backup) if backup.exists() else "",
    }


def metadata(results_dir: str | Path) -> dict[str, Any]:
    """Read V13 best-memory metadata, or return an empty mapping."""
    source = best_file(results_dir)
    if not source.exists():
        return {}
    try:
        df = pd.read_excel(source, sheet_name="v13_metadata")
        if df.empty:
            return {}
        return {str(k): v for k, v in df.iloc[0].to_dict().items()}
    except Exception:
        return {}
