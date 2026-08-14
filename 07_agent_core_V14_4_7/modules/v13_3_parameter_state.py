from __future__ import annotations

"""Dedicated V13.3 best-parameter memory."""

from datetime import datetime
from pathlib import Path
from typing import Any
import shutil
import pandas as pd
from openpyxl import load_workbook


BEST_FILE_NAME = "best_parameters_V13_3.xlsx"


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def best_file(results_dir: str | Path) -> Path:
    return Path(results_dir) / BEST_FILE_NAME


def has_best(results_dir: str | Path) -> bool:
    return best_file(results_dir).exists()


def save_best(
    config_file: str | Path,
    results_dir: str | Path,
    *,
    total_score: float,
    run_folder: str,
    objective_reference_file: str = "",
    objective_mode: str = "multi_species_frozen_reference_weighted_rmse",
    optimizer_version: str = "V13.3",
    stage: str = "",
    diagnostics: dict[str, Any] | None = None,
) -> dict[str, Any]:
    config_file = Path(config_file)
    output = best_file(results_dir)
    output.parent.mkdir(parents=True, exist_ok=True)
    if not config_file.exists():
        raise FileNotFoundError(config_file)

    backup = None
    if output.exists():
        backup = output.with_name(f"{output.stem}_previous_{datetime.now():%Y%m%d_%H%M%S}{output.suffix}")
        shutil.copy2(output, backup)

    shutil.copy2(config_file, output)
    metadata = {
        "created_at": _now(),
        "source_run_folder": str(run_folder),
        "TOTAL_SCORE": float(total_score),
        "objective_reference_file": str(objective_reference_file),
        "objective_mode": str(objective_mode),
        "optimizer_version": str(optimizer_version),
        "stage": str(stage),
        **(diagnostics or {}),
    }

    workbook = load_workbook(output)
    if "v13_3_metadata" in workbook.sheetnames:
        del workbook["v13_3_metadata"]
    ws = workbook.create_sheet("v13_3_metadata")
    ws.append(list(metadata.keys()))
    ws.append(list(metadata.values()))
    workbook.save(output)

    return {"action": "saved_v13_3_best", "best_file": str(output), "backup_file": str(backup or ""), **metadata}


def restore_best(config_file: str | Path, results_dir: str | Path) -> dict[str, Any]:
    config_file = Path(config_file)
    source = best_file(results_dir)
    if not source.exists():
        raise FileNotFoundError(source)

    backup_dir = Path(results_dir) / "v13_3_parameter_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"{config_file.stem}_before_v13_3_restore_{datetime.now():%Y%m%d_%H%M%S}{config_file.suffix}"
    if config_file.exists():
        shutil.copy2(config_file, backup)
    shutil.copy2(source, config_file)
    return {"action": "restored_v13_3_best", "best_file": str(source), "config_backup_file": str(backup)}


def metadata(results_dir: str | Path) -> dict[str, Any]:
    source = best_file(results_dir)
    if not source.exists():
        return {}
    try:
        df = pd.read_excel(source, sheet_name="v13_3_metadata")
        return {} if df.empty else df.iloc[0].to_dict()
    except Exception:
        return {}
